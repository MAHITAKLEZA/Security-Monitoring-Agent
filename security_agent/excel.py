"""Excel (.xlsx) version of a scan report: Summary, Findings and Recommendations sheets."""

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEVERITY_FILL = {
    "CRITICAL": "F8D7DA", "HIGH": "FCE4D6", "MEDIUM": "FFF2CC", "LOW": "DDEBF7", "INFO": "EDEDED",
}
CHECK_LABELS = {
    "ssl": "SSL Monitoring", "headers": "Security Headers", "exposed_files": "Exposed Files",
    "changes": "Suspicious Changes", "malware": "Malware Detection", "vulnerabilities": "Vulnerability Check",
    "dns": "DNS & Domain", "third_party": "3rd Party Resources", "wordpress": "WordPress", "availability": "Availability", "agent": "Agent",
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _scan_time(iso):
    try:
        return datetime.fromisoformat(iso).astimezone().strftime("%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return iso or ""


def _table(ws, start_row, headers, rows, widths, severity_col=None):
    """Write a formatted table; returns the next free row."""
    for col, (title, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(start_row, col, title)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = max(ws.column_dimensions[get_column_letter(col)].width or 0, width)
    for r, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(r, col, value)
            cell.border, cell.alignment = BORDER, WRAP
        if severity_col is not None:
            sev = str(row[severity_col])
            if sev in SEVERITY_FILL:
                cell = ws.cell(r, severity_col + 1)
                cell.fill = PatternFill("solid", fgColor=SEVERITY_FILL[sev])
                cell.font = Font(bold=True)
    return start_row + len(rows) + 1


def to_xlsx(report):
    """Build the workbook for one site's report and return it as bytes."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Security Report: {report.get('site') or report['target']}"
    ws["A1"].font = Font(bold=True, size=14)
    info = [("Website", report["target"]), ("Scan time", _scan_time(report["started_at"])),
            ("Security status", report["security_status"]), ("Risk score", f"{report['risk_score']}/100")]
    for i, (label, value) in enumerate(info, 3):
        ws.cell(i, 1, label).font = Font(bold=True)
        ws.cell(i, 2, value)
    row = _table(ws, 8, ["Severity", "Count"], list(report["severity_counts"].items()), [22, 14], severity_col=0)
    _table(ws, row + 1, ["Check", "Findings", "Worst severity"],
           [(CHECK_LABELS.get(name, name), c["findings"], c["worst"]) for name, c in report["checks"].items()],
           [22, 14, 16], severity_col=2)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width, len(report["target"]) + 4)

    ws = wb.create_sheet("Findings")
    rows = [(f["severity"], CHECK_LABELS.get(f["check"], f["check"]), f["title"], f["detail"], f["recommendation"],
             json.dumps(f["evidence"], default=str)[:2000] if f.get("evidence") else "")
            for f in report["findings"]]
    _table(ws, 1, ["Severity", "Check", "Finding", "Detail", "Recommendation", "Evidence"], rows,
           [12, 20, 45, 50, 55, 45], severity_col=0)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"

    ws = wb.create_sheet("Recommendations")
    rows = [(i, r["severity"], r["title"], r["recommendation"]) for i, r in enumerate(report["recommendations"], 1)]
    _table(ws, 1, ["#", "Severity", "Issue", "Recommendation"], rows, [5, 12, 50, 80], severity_col=1)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
